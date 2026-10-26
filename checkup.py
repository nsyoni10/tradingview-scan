"""
checkup.py — סקיל checkup: דוח לונג יומי לרשימת "Check list" מטריידינגוויו.

ציון כניסה ללונג בתחילת מהלך אחרי תיקון: כיוון (DI/ADX), מרובוזו, ווליום, CCI, SMA150.
הפלט הוא xlsx שמועלה ל-Google Drive ומומר ל-Google Sheet בשם "Summary of <date>".

קלט JSON כמו ב-ticker_report.py, ובנוסף לכל טיקר: "SMA150", "RVOL", "CCI", "CCI1", "ADX1" (ערכי אתמול),
ו-"bars": 12 נרות יומיים אחרונים [o,h,l,c,v] מהישן לחדש — מהם נגזרים prev/last, RVOL (יום אחרון מול 10 שלפניו) ומרובוזו.
שימוש:  python3 checkup.py checkup_data/data_YYYY-MM-DD.json out.xlsx
"""
import csv
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ticker_report import candle, levels, trend

RVOL_MIN = 1.2
MARU_BODY = 0.85      # מרובוזו: גוף ≥ 85% מטווח הנר
MARU_MIN_ATR = 0.5    # ...וגוף ≥ 0.5 ATR (נר עם עוצמה, לא נר קטנטן)
MARU_DAYS = 3         # מחפשים מרובוזו ב-3 ימי המסחר האחרונים
OK, BAD = "✅", "❌"


# ── ווליום: יום אחרון מול ממוצע 10 הימים שלפניו (בלי היום עצמו) ────────
def rvol(d):
    bars = d.get("bars")
    if bars and len(bars) >= 11:
        prior = [b[4] for b in bars[-11:-1]]
        return bars[-1][4] / (sum(prior) / len(prior))
    return d["RVOL"]


# ── מרובוזו ירוק ב-N הימים האחרונים ─────────────────────────────────────
def marubozu(d):
    bars = d.get("bars") or [d["prev"] + [0], d["last"] + [0]]
    for ago, (o, h, l, c, _) in enumerate(reversed(bars[-MARU_DAYS:])):
        rng = h - l
        if rng > 0 and c > o and (c - o) / rng >= MARU_BODY and (c - o) >= MARU_MIN_ATR * d["ATR"]:
            return ago
    return None


# ── ציון כניסה ללונג 0-100 (תחילת מהלך אחרי תיקון, דגש על כיוון) ─────────
def score(d, cd, res, rv, maru):
    pts = {}
    c, o = d["last"][3], d["last"][0]
    adx, dip, dim = d["ADX"], d["DIp"], d["DIm"]

    # כיוון (30): +DI מעל -DI, ADX סביב 20 ומעלה, ADX עולה — נספר רק כשהכיוון ללונג
    if dip > dim:
        pts["כיוון"] = 15 + (10 if adx >= 20 else 5 if adx >= 18 else 0) + (5 if adx > d.get("ADX1", adx) else 0)
    else:
        pts["כיוון"] = 0

    # נר עם עוצמה (25): מרובוזו ירוק ב-3 הימים האחרונים + דפוס הנר האחרון
    pts["נר"] = (15 if maru == 0 else 10 if maru is not None else 0) + cd["pts"] * 0.4

    # ווליום (15): רק על נר ירוק — ווליום גבוה על נר אדום הוא מכירה
    if c > o:
        pts["ווליום"] = 15 if rv >= 1.5 else 10 if rv >= RVOL_MIN else 4 if rv >= 1.0 else 0
    else:
        pts["ווליום"] = 0

    # CCI (15): יציאה מתיקון — חוצה למעלה את -100 או את 0, או עולה מתחת ל-100
    cci, cci1 = d["CCI"], d.get("CCI1", d["CCI"])
    if cci > cci1 and ((cci1 < -100 <= cci) or (cci1 < 0 <= cci)):
        pts["CCI"] = 15
    elif cci > cci1 and cci < 100:
        pts["CCI"] = 10
    elif cci > cci1 and cci < 200:
        pts["CCI"] = 4
    else:
        pts["CCI"] = 0

    # מעל SMA150 (10)
    pts["SMA150"] = 10 if c > d["SMA150"] else 0

    # לא מתוח (5): עד 1.5 ATR מעל EMA20 — אחרת המהלך כבר קרה
    pts["לא מתוח"] = 5 if (c - d["EMA20"]) / d["ATR"] <= 1.5 else 0

    room = (res[0][0] - c) / d["ATR"] if res else 3
    return round(sum(pts.values())), room


def build(data, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.rightToLeft = True

    headers = [
        "טיקר", "ציון כניסה", "מעל SMA150", "RVOL ≥ 1.2", "כיוון", "סגירה", "שינוי %",
        "SMA150", "RVOL", "ATR", "ATR %", "ADX", "+DI", "-DI", "CCI", "EMA20", "EMA50",
        "צבע נר", "מרובוזו (3 ימים)", "דפוס נר",
        "התנגדות 1", "מקור R1", "מרחק R1 (ATR)", "התנגדות 2",
        "תמיכה 1", "מקור S1", "מרחק S1 (ATR)", "תמיכה 2",
        "נפח ממוצע 10י", "הערות",
    ]
    ws.append(headers)

    rows = []
    for sym, d in data["tickers"].items():
        if d.get("bars"):
            d["prev"], d["last"] = d["bars"][-2][:4], d["bars"][-1][:4]
        d["RVOL"] = rvol(d)
        maru = marubozu(d)
        tr = trend(d)
        cd = candle(d, tr)
        res, sup = levels(d)
        sc, room = score(d, cd, res, d["RVOL"], maru)
        c, pc = d["last"][3], d["prev"][3]
        above = c > d["SMA150"]
        rvol_ok = d["RVOL"] >= RVOL_MIN

        notes = []
        if d["ATR"] / c < 0.005:
            notes.append("תנודתיות כמעט אפסית — מחיר 'קפוא'. לא לסחור טכנית")
            sc = 0
        if d["DIm"] > d["DIp"]:
            notes.append("-DI מעל +DI — הכיוון עדיין לא ללונג")
        if d["RVOL"] >= RVOL_MIN and d["last"][3] < d["last"][0]:
            notes.append("ווליום גבוה על נר אדום — מכירה")
        if (d["last"][3] - d["EMA20"]) / d["ATR"] > 1.5:
            notes.append("מתוח מעל EMA20 — ייתכן שהמהלך כבר קרה")

        def lvl(lst, i):
            return (lst[i][0], lst[i][1]) if len(lst) > i else (None, "")

        r1, r1n = lvl(res, 0)
        r2, _ = lvl(res, 1)
        s1, s1n = lvl(sup, 0)
        s2, _ = lvl(sup, 1)
        rows.append([
            sym.split(":")[1], sc, OK if above else BAD, OK if rvol_ok else BAD,
            ("▲ +DI" if d["DIp"] > d["DIm"] else "▼ -DI") + f" | ADX {d['ADX']:.0f}" + (" עולה" if d["ADX"] > d.get("ADX1", d["ADX"]) else " יורד"),
            c, (c / pc - 1) * 100, d["SMA150"], d["RVOL"],
            d["ATR"], d["ATR"] / c * 100, d["ADX"], d["DIp"], d["DIm"], d["CCI"],
            d["EMA20"], d["EMA50"], cd["color"],
            "—" if maru is None else ("היום" if maru == 0 else f"לפני {maru} ימים" if maru > 1 else "אתמול"),
            cd["tags"],
            r1, r1n, (r1 - c) / d["ATR"] if r1 else None, r2,
            s1, s1n, (c - s1) / d["ATR"] if s1 else None, s2,
            d["vol10"], "; ".join(notes),
        ])

    rows.sort(key=lambda r: -r[1])
    for r in rows:
        ws.append(r)

    # עיצוב
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"

    fmt = {
        6: "0.00", 7: "+0.00;-0.00", 8: "0.00", 9: "0.00", 10: "0.00", 11: "0.0",
        12: "0.0", 13: "0.0", 14: "0.0", 15: "0", 16: "0.00", 17: "0.00",
        21: "0.00", 23: "0.0", 24: "0.00", 25: "0.00", 27: "0.0", 28: "0.00", 29: "#,##0",
    }
    green, red = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.column in fmt:
                cell.number_format = fmt[cell.column]
        row[0].font = Font(bold=True)
        row[1].font = Font(bold=True, size=12)
        for i in (2, 3):
            row[i].fill = green if row[i].value == OK else red
        row[4].fill = green if "▲" in row[4].value else red
        row[18].fill = green if row[18].value != "—" else PatternFill()
        row[17].fill = green if row[17].value == "ירוק" else red if row[17].value == "אדום" else PatternFill()
        row[19].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        row[29].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    last = ws.max_row
    ws.conditional_formatting.add(
        f"B2:B{last}",
        ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                       mid_type="num", mid_value=50, mid_color="FFEB84",
                       end_type="num", end_value=100, end_color="63BE7B"))

    widths = {1: 8, 2: 9, 3: 9, 4: 9, 5: 16, 19: 12, 20: 30, 22: 18, 26: 18, 30: 40}
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 10)

    # גיליון הסבר
    ex = wb.create_sheet("הסבר")
    ex.sheet_view.rightToLeft = True
    lines = [
        ("מקור", f"רשימת המעקב '{data.get('watchlist', 'Check list')}' בטריידינגוויו, גרף יומי. נר אחרון: {data.get('bar_date', data['date'])}"),
        ("מעל SMA150", "סגירה מעל הממוצע הנע הפשוט של 150 ימים = ✅"),
        ("RVOL ≥ 1.2", "ווליום יום המסחר האחרון שנסגר חלקי ממוצע 10 הימים שלפניו (בלי היום עצמו), מחושב מהנרות. "
                       "1.2 ומעלה = ✅. בלי פרה/אפטר-מרקט"),
        ("כיוון", "+DI מול -DI, ערך ADX(14), והאם ADX עלה מאתמול. ADX סביב 20 ומעלה עם +DI למעלה = כיוון לונג"),
        ("CCI", "CCI(20) יומי. מחפשים יציאה מתיקון: CCI עולה, חוצה למעלה את -100 או את 0"),
        ("מרובוזו", "נר ירוק עם גוף ≥ 85% מהטווח וגוף ≥ 0.5 ATR, באחד מ-3 ימי המסחר האחרונים"),
        ("פטיש / איש תלוי", "צל תחתון ≥ 2× גוף, צל עליון ≤ 30% מהטווח. במגמה עולה = איש תלוי, אחרת = פטיש"),
        ("נר 2/3", "נר ירוק שגופו ≥ 2/3 מטווח הנר, אחרי נר אדום"),
        ("תמיכה/התנגדות", "Pivot חודשי, שיא/שפל חודש ו-3 חודשים, EMA20/50, גבוה/נמוך של 2 הנרות האחרונים. "
                          "רמות במרחק < 0.25 ATR מאוחדות"),
        ("ציון כניסה 0-100", "כיוון 30 (+DI>-DI 15, ADX≥20 10 / 18-20 5, ADX עולה 5 — רק כש-+DI>-DI) | "
                             "נר עם עוצמה 25 (מרובוזו היום 15 / ב-3 ימים 10 + דפוס נר אחרון עד 10) | "
                             "ווליום על נר ירוק 15 | CCI יוצא מתיקון 15 | מעל SMA150 10 | לא מתוח מעל EMA20 5"),
        ("שים לב", "הציון הוא כלי מיון, לא המלצת השקעה"),
    ]
    for k, v in lines:
        ex.append([k, v])
    ex.column_dimensions["A"].width = 18
    ex.column_dimensions["B"].width = 110
    for row in ex.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, horizontal="right")

    wb.save(out)
    write_csv(headers, rows, out.rsplit(".", 1)[0] + ".csv")
    return rows


def write_csv(headers, rows, path):
    """CSV להעלאה ל-Google Drive (מומר אוטומטית ל-Google Sheet)."""
    def fmt(v):
        if isinstance(v, float):
            return round(v, 2)
        return "" if v is None else v

    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in rows:
            w.writerow([fmt(v) for v in r])


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    with open(src, encoding="utf-8") as f:
        rows = build(json.load(f), dst)
    for r in rows:
        print(f"{r[0]:6} {r[1]:3} SMA150{r[2]} RVOL{r[3]} {r[4]:22} RVOL={r[8]:.2f} CCI={r[14]:7.1f} מרובוזו={r[18]} | {r[19]}")
