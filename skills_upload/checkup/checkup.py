"""
checkup.py — סקיל checkup: דוח לונג יומי לרשימת "Check list" מטריידינגוויו.

מוסיף על ticker_report.py: מעל SMA150 יומי, ווליום יחסי (RVOL) ≥ 1.2.
הפלט הוא xlsx שמועלה ל-Google Drive ומומר ל-Google Sheet בשם "Summary of <date>".

קלט JSON כמו ב-ticker_report.py, ובנוסף לכל טיקר: "SMA150", "RVOL", ואופציונלית "last_vol" (ווליום הבר האחרון; אם קיים — RVOL מחושב = last_vol / vol10).
שימוש:  python3 checkup.py checkup_data/data_YYYY-MM-DD.json out.xlsx
"""
import csv
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ticker_report import candle, levels, trend

RVOL_MIN = 1.2
OK, BAD = "✅", "❌"


# ── ציון לונג 0-100 ───────────────────────────────────────────────────
def score(d, cd, res):
    pts = {}
    c, e20, e50 = d["last"][3], d["EMA20"], d["EMA50"]
    adx, dip, dim = d["ADX"], d["DIp"], d["DIm"]

    # מבנה EMA (20)
    pts["EMA"] = (8 if c > e20 else 0) + (8 if e20 > e50 else 0) + (4 if c > e50 else 0)

    # מעל SMA150 (10)
    pts["SMA150"] = 10 if c > d["SMA150"] else 0

    # ADX בכיוון הלונג (15)
    if dip > dim:
        pts["ADX"] = 15 if adx >= 25 else 10 if adx >= 20 else 5
    else:
        pts["ADX"] = 0 if adx >= 25 else 3

    # עליונות +DI (10)
    pts["DI"] = max(min((dip - dim) / 20 * 10, 10), 0)

    # נר אחרון (20) — ממיר את סקאלת 0-25 של ticker_report
    pts["נר"] = cd["pts"] * 0.8

    # RVOL (10)
    pts["RVOL"] = 10 if d["RVOL"] >= RVOL_MIN else 5 if d["RVOL"] >= 1.0 else 0

    # RSI (10)
    rsi = d["RSI"]
    pts["RSI"] = 10 if 45 <= rsi <= 68 else 5 if 35 <= rsi < 45 or 68 < rsi <= 75 else 0

    # מרחב עד ההתנגדות (5)
    room = (res[0][0] - c) / d["ATR"] if res else 3
    pts["מרחב"] = 5 if room >= 1.5 else 2 if room >= 0.75 else 0

    return round(sum(pts.values())), room


def build(data, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.rightToLeft = True

    headers = [
        "טיקר", "ציון לונג", "מעל SMA150", "RVOL ≥ 1.2", "מגמה", "סגירה", "שינוי %",
        "SMA150", "RVOL", "ATR", "ATR %", "ADX", "+DI", "-DI", "RSI", "EMA20", "EMA50",
        "צבע נר", "גוף/טווח", "דפוס נר",
        "התנגדות 1", "מקור R1", "מרחק R1 (ATR)", "התנגדות 2",
        "תמיכה 1", "מקור S1", "מרחק S1 (ATR)", "תמיכה 2",
        "נפח ממוצע 10י", "הערות",
    ]
    ws.append(headers)

    rows = []
    for sym, d in data["tickers"].items():
        rvol_note = None
        if d.get("last_vol") and d.get("vol10"):
            own = d["last_vol"] / d["vol10"]
            if d.get("RVOL") and abs(own - d["RVOL"]) / d["RVOL"] > 0.15:
                rvol_note = f"RVOL של טריידינגוויו ({d['RVOL']:.2f}) שונה מהחישוב לפי הנר האחרון ({own:.2f})"
            d["RVOL"] = own
        tr = trend(d)
        cd = candle(d, tr)
        res, sup = levels(d)
        sc, room = score(d, cd, res)
        c, pc = d["last"][3], d["prev"][3]
        above = c > d["SMA150"]
        rvol_ok = d["RVOL"] >= RVOL_MIN

        notes = [rvol_note] if rvol_note else []
        if d["ATR"] / c < 0.005:
            notes.append("תנודתיות כמעט אפסית — מחיר 'קפוא'. לא לסחור טכנית")
            sc = 0
        if tr == "DOWN":
            notes.append("מגמת ירידה — לונג נגד המגמה")
        if d["RSI"] > 70:
            notes.append("RSI קנוי-יתר")
        if d["RSI"] < 30:
            notes.append("RSI מכור-יתר")

        def lvl(lst, i):
            return (lst[i][0], lst[i][1]) if len(lst) > i else (None, "")

        r1, r1n = lvl(res, 0)
        r2, _ = lvl(res, 1)
        s1, s1n = lvl(sup, 0)
        s2, _ = lvl(sup, 1)
        rows.append([
            sym.split(":")[1], sc, OK if above else BAD, OK if rvol_ok else BAD,
            {"UP": "עולה ▲", "DOWN": "יורדת ▼", "MIXED": "מעורבת"}[tr],
            c, (c / pc - 1) * 100, d["SMA150"], d["RVOL"],
            d["ATR"], d["ATR"] / c * 100, d["ADX"], d["DIp"], d["DIm"], d["RSI"],
            d["EMA20"], d["EMA50"], cd["color"], cd["body_pct"], cd["tags"],
            r1, r1n, (r1 - c) / d["ATR"] if r1 else None, r2,
            s1, s1n, (c - s1) / d["ATR"] if s1 else None, s2,
            d["vol10"], "; ".join(notes),
        ])

    rows.sort(key=lambda r: -r[1])
    for r in rows:
        ws.append(r)

    # עיצוב
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"

    fmt = {
        6: "0.00", 7: "+0.00;-0.00", 8: "0.00", 9: "0.00", 10: "0.00", 11: "0.0",
        12: "0.0", 13: "0.0", 14: "0.0", 15: "0.0", 16: "0.00", 17: "0.00", 19: "0%",
        21: "0.00", 23: "0.0", 24: "0.00", 25: "0.00", 27: "0.0", 28: "0.00", 29: "#,##0",
    }
    green, red = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.column in fmt:
                cell.number_format = fmt[cell.column]
        row[0].font = Font(bold=True)
        row[1].font = Font(bold=True, size=12)
        for i in (2, 3):
            row[i].fill = green if row[i].value == OK else red
        row[4].fill = green if "עולה" in row[4].value else red if "יורדת" in row[4].value else PatternFill()
        row[17].fill = green if row[17].value == "ירוק" else red if row[17].value == "אדום" else PatternFill()
        row[19].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        row[29].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    last = ws.max_row
    ws.conditional_formatting.add(
        f"B2:B{last}",
        ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                       mid_type="num", mid_value=50, mid_color="FFEB84",
                       end_type="num", end_value=100, end_color="63BE7B"))

    widths = {1: 8, 2: 9, 3: 9, 4: 9, 5: 9, 20: 30, 22: 18, 26: 18, 30: 40}
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 10)

    # גיליון הסבר
    ex = wb.create_sheet("הסבר")
    ex.sheet_view.rightToLeft = True
    lines = [
        ("מקור", f"רשימת המעקב '{data.get('watchlist', 'Check list')}' בטריידינגוויו, גרף יומי. נר אחרון: {data.get('bar_date', data['date'])}"),
        ("מעל SMA150", "סגירה מעל הממוצע הנע הפשוט של 150 ימים = ✅"),
        ("RVOL ≥ 1.2", "ווליום יום המסחר האחרון שנסגר חלקי ממוצע 10 ימים. 1.2 ומעלה = ✅. בלי פרה/אפטר-מרקט"),
        ("ATR / ADX", "ATR(14) ו-ADX(14) יומיים. ADX ≥ 25 = מגמה חזקה, 20-25 = מתפתחת, < 20 = אין מגמה"),
        ("מגמה", "עולה: סגירה > EMA20 > EMA50 ו-+DI > -DI. יורדת: ההפך. אחרת: מעורבת"),
        ("פטיש / איש תלוי", "צל תחתון ≥ 2× גוף, צל עליון ≤ 30% מהטווח. במגמה עולה = איש תלוי, אחרת = פטיש"),
        ("נר 2/3", "נר ירוק שגופו ≥ 2/3 מטווח הנר, אחרי נר אדום"),
        ("תמיכה/התנגדות", "Pivot חודשי, שיא/שפל חודש ו-3 חודשים, EMA20/50, גבוה/נמוך של 2 הנרות האחרונים. "
                          "רמות במרחק < 0.25 ATR מאוחדות"),
        ("ציון לונג 0-100", "מבנה EMA 20 | מעל SMA150 10 | ADX עם +DI>-DI 15 | עליונות +DI 10 | נר 20 | "
                            "RVOL 10 | RSI 45-68 10 | מרחב להתנגדות 5"),
        ("שים לב", "הציון הוא כלי מיון, לא המלצת השקעה"),
    ]
    for k, v in lines:
        ex.append([k, v])
    ex.column_dimensions["A"].width = 18
    ex.column_dimensions["B"].width = 110
    for row in ex.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, horizontal="right")

    wb.save(out)
    write_csv(headers, rows, out.rsplit(".", 1)[0] + ".csv")
    return rows


def write_csv(headers, rows, path):
    """CSV להעלאה ל-Google Drive (מומר אוטומטית ל-Google Sheet)."""
    def fmt(v):
        if isinstance(v, float):
            return round(v, 2)
        return "" if v is None else v

    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in rows:
            w.writerow([fmt(v) for v in r])


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    with open(src, encoding="utf-8") as f:
        rows = build(json.load(f), dst)
    for r in rows:
        print(f"{r[0]:6} {r[1]:3} SMA150{r[2]} RVOL{r[3]} {r[4]:8} RVOL={r[8]:.2f} ADX={r[11]:5.1f} | {r[19]}")
